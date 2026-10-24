"""Build the Word version of the paper from paper/ieee_paper.html.

Both outputs come from the one HTML source, so the .docx and the .pdf cannot
drift apart: edit the HTML, run this and build_pdf.py, and the two agree.

The IEEE two-column layout is done the way the official Word template does it —
a continuous section break down to one column around every full-width figure or
table, and straight back to two afterwards. Word balances the columns at each
break, which is what puts a spanning figure across the page.

Figures are placed as PNG because python-docx cannot embed SVG; the rasters are
2064px across, which is 288 dpi at the 7.17in text width.

    python paper/preview_svg.py <each svg>      # refresh the rasters
    python paper/build_docx.py
"""
from __future__ import annotations

import re
from pathlib import Path

from bs4 import BeautifulSoup, NavigableString, Tag
from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
from docx.oxml import OxmlElement
from docx.shared import Inches, Pt, RGBColor

HERE = Path(__file__).resolve().parent
SRC = HERE / "ieee_paper.html"
FIGS = HERE / "_docx_figs"
OUT = HERE.parent / "TeamSync_IEEE_Paper.docx"

SERIF = "Times New Roman"
MONO = "Courier New"
BODY_PT = Pt(10)
SMALL_PT = Pt(8)

TEXT_WIDTH = 7.16          # inches, letter minus 0.67in side margins
COL_WIDTH = 3.45

# column widths, in inches, per table in document order
TABLE_WIDTHS = [
    [1.06, 1.02, 0.94, 0.43],                                         # I (claims)
    [0.85, 1.90, 0.70],                                               # II
    [0.80, 1.25, 0.75, 0.65],                                         # III
    [0.64, 1.22, 0.60, 0.48, 0.44, 0.50, 0.42, 1.56, 1.10],           # IV (wide)
    [1.36, 0.42, 0.42, 0.42, 0.42, 0.42],                             # V
]


# ----------------------------------------------------------------- plumbing
def set_columns(section, num: int, space_in: float = 0.25) -> None:
    sect_pr = section._sectPr
    cols = sect_pr.find(qn("w:cols"))
    if cols is None:
        cols = OxmlElement("w:cols")
        sect_pr.append(cols)
    cols.set(qn("w:num"), str(num))
    cols.set(qn("w:space"), str(int(space_in * 1440)))
    cols.set(qn("w:equalWidth"), "1")


def border(cell, edge: str, sz: int) -> None:
    """sz is eighths of a point: 8 = 1pt."""
    tc_pr = cell._tc.get_or_add_tcPr()
    borders = tc_pr.find(qn("w:tcBorders"))
    if borders is None:
        borders = OxmlElement("w:tcBorders")
        tc_pr.append(borders)
    el = borders.find(qn(f"w:{edge}"))
    if el is None:
        el = OxmlElement(f"w:{edge}")
        borders.append(el)
    el.set(qn("w:val"), "single")
    el.set(qn("w:sz"), str(sz))
    el.set(qn("w:color"), "000000")


def no_table_borders(table) -> None:
    tbl_pr = table._tbl.tblPr
    borders = OxmlElement("w:tblBorders")
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        el = OxmlElement(f"w:{edge}")
        el.set(qn("w:val"), "none")
        el.set(qn("w:sz"), "0")
        borders.append(el)
    tbl_pr.append(borders)


def set_cell_margins(table, left=36, right=36, top=12, bottom=12) -> None:
    """Twips. Word's 0.08in default side padding is too generous at 8pt."""
    mar = OxmlElement("w:tblCellMar")
    for edge, val in (("top", top), ("left", left), ("bottom", bottom), ("right", right)):
        el = OxmlElement(f"w:{edge}")
        el.set(qn("w:w"), str(val))
        el.set(qn("w:type"), "dxa")
        mar.append(el)
    table._tbl.tblPr.append(mar)


def keep_table_whole(table) -> None:
    """Stop Word splitting a table (or its caption) over a page break."""
    rows = table.rows
    for i, row in enumerate(rows):
        tr_pr = row._tr.get_or_add_trPr()
        cant = OxmlElement("w:cantSplit")
        tr_pr.append(cant)
        if i < len(rows) - 1:
            for cell in row.cells:
                for par in cell.paragraphs:
                    par.paragraph_format.keep_with_next = True


BODY_LEAD = Pt(11.4)     # exact leading for 10pt Times
SMALL_LEAD = Pt(9.0)     # exact leading for 8pt Times


def tight(par, before=0, after=0, lead=BODY_LEAD) -> None:
    pf = par.paragraph_format
    pf.space_before = Pt(before)
    pf.space_after = Pt(after)
    pf.line_spacing = lead


# ----------------------------------------------------------------- inline
def add_runs(par, node, *, size=BODY_PT, italic=False, bold=False) -> None:
    for child in node.children:
        if isinstance(child, NavigableString):
            txt = re.sub(r"\s+", " ", str(child))
            if not txt:
                continue
            run = par.add_run(txt)
            run.font.name = SERIF
            run.font.size = size
            run.italic = italic
            run.bold = bold
        elif isinstance(child, Tag):
            if child.name in ("em", "i", "h4"):
                add_runs(par, child, size=size, italic=True, bold=bold)
            elif child.name in ("strong", "b"):
                add_runs(par, child, size=size, italic=italic, bold=True)
            elif child.name == "span" and "mono" in (child.get("class") or []):
                run = par.add_run(re.sub(r"\s+", " ", child.get_text()))
                run.font.name = MONO
                run.font.size = Pt(size.pt - 1.2)
            elif child.name == "span" and "best" in (child.get("class") or []):
                add_runs(par, child, size=size, italic=italic, bold=True)
            elif child.name == "br":
                par.add_run().add_break()
            else:
                add_runs(par, child, size=size, italic=italic, bold=bold)


# ----------------------------------------------------------------- blocks
def emit_paragraph(doc, el) -> None:
    classes = el.get("class") or []
    par = doc.add_paragraph()
    par.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    tight(par)
    if "abstract" in classes or "keywords" in classes:
        par.paragraph_format.line_spacing = Pt(10.2)
        add_runs(par, el, size=Pt(9), bold=True)
        par.paragraph_format.space_after = Pt(6)
    else:
        if "first" not in classes:
            par.paragraph_format.first_line_indent = Inches(0.17)
        add_runs(par, el)


def emit_heading(doc, el, level: int) -> None:
    par = doc.add_paragraph()
    tight(par, before=10 if level == 2 else 7, after=3)
    par.paragraph_format.keep_with_next = True
    if level == 2:
        par.alignment = WD_ALIGN_PARAGRAPH.CENTER
        run = par.add_run(el.get_text().replace(" ", " ").strip())
        run.font.name = SERIF
        run.font.size = BODY_PT
        run.font.small_caps = True
    else:
        par.alignment = WD_ALIGN_PARAGRAPH.LEFT
        run = par.add_run(el.get_text().strip())
        run.font.name = SERIF
        run.font.size = BODY_PT
        run.italic = True


def emit_list(doc, el) -> None:
    for li in el.find_all("li", recursive=False):
        par = doc.add_paragraph(style="List Bullet")
        par.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
        tight(par, after=2)
        par.paragraph_format.left_indent = Inches(0.22)
        par.paragraph_format.first_line_indent = Inches(-0.12)
        add_runs(par, li)


def emit_figure(doc, el, width_in: float) -> None:
    img = el.find("img")
    name = Path(img["src"]).with_suffix(".png").name
    path = FIGS / name
    assert path.exists(), f"missing raster: {path}"

    par = doc.add_paragraph()
    par.alignment = WD_ALIGN_PARAGRAPH.CENTER
    tight(par, before=6, after=3)
    # exact leading crops an inline image to the line box; a float makes Word
    # write lineRule="auto", which grows the line to the picture
    par.paragraph_format.line_spacing = 1.0
    par.add_run().add_picture(str(path), width=Inches(width_in))

    cap = el.find("figcaption")
    if cap:
        p = doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
        tight(p, after=8, lead=SMALL_LEAD)
        add_runs(p, cap, size=SMALL_PT)


def emit_table(doc, wrap, widths: list[float]) -> None:
    cap = wrap.find("p", class_="tcap")
    if cap:
        for i, line in enumerate(cap.stripped_strings):
            p = doc.add_paragraph()
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            tight(p, before=6 if i == 0 else 0, after=0 if i == 0 else 2, lead=SMALL_LEAD)
            p.paragraph_format.keep_with_next = True
            run = p.add_run(line if i else line.upper())
            run.font.name = SERIF
            run.font.size = SMALL_PT

    html_tbl = wrap.find("table")
    head = html_tbl.find("thead").find_all("tr")
    body = html_tbl.find("tbody").find_all("tr")
    rows = head + body
    ncols = len(widths)

    table = doc.add_table(rows=len(rows), cols=ncols)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    no_table_borders(table)
    set_cell_margins(table)
    for col, w in zip(table.columns, widths):
        col.width = Inches(w)

    for r, tr in enumerate(rows):
        cells = tr.find_all(["th", "td"])
        is_head = tr in head
        for c, (cell, html_cell) in enumerate(zip(table.rows[r].cells, cells)):
            cell.width = Inches(widths[c])
            par = cell.paragraphs[0]
            tight(par, before=1, after=1, lead=SMALL_LEAD)
            klass = html_cell.get("class") or []
            if "n" in klass:
                par.alignment = WD_ALIGN_PARAGRAPH.RIGHT
            elif "c" in klass:
                par.alignment = WD_ALIGN_PARAGRAPH.CENTER
            else:
                par.alignment = WD_ALIGN_PARAGRAPH.LEFT
            add_runs(par, html_cell, size=SMALL_PT, bold=is_head)

            if is_head and r == 0:
                border(cell, "top", 10)
            if is_head and tr is head[-1]:
                border(cell, "bottom", 6)
            if r == len(rows) - 1:
                border(cell, "bottom", 10)
            if "rule" in (tr.get("class") or []):
                border(cell, "top", 4)

    keep_table_whole(table)
    doc.add_paragraph().paragraph_format.space_after = Pt(4)


def emit_references(doc, ol) -> None:
    for i, li in enumerate(ol.find_all("li", recursive=False), 1):
        par = doc.add_paragraph()
        par.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
        tight(par, after=1.5, lead=SMALL_LEAD)
        par.paragraph_format.left_indent = Inches(0.21)
        par.paragraph_format.first_line_indent = Inches(-0.21)
        run = par.add_run(f"[{i}] ")
        run.font.name = SERIF
        run.font.size = SMALL_PT
        add_runs(par, li, size=SMALL_PT)


# ----------------------------------------------------------------- document
def main() -> None:
    soup = BeautifulSoup(SRC.read_text(encoding="utf-8"), "html.parser")
    doc = Document()

    normal = doc.styles["Normal"]
    normal.font.name = SERIF
    normal.font.size = BODY_PT
    normal.element.rPr.rFonts.set(qn("w:eastAsia"), SERIF)
    normal.paragraph_format.space_after = Pt(0)
    normal.paragraph_format.line_spacing = Pt(11.4)

    first = doc.sections[0]
    first.page_width, first.page_height = Inches(8.5), Inches(11)
    first.top_margin, first.bottom_margin = Inches(0.75), Inches(1.0)
    first.left_margin = first.right_margin = Inches(0.67)
    set_columns(first, 1)

    # --- title block, full width
    tb = soup.find("div", class_="titleblock")
    h1 = tb.find("h1")
    par = doc.add_paragraph()
    par.alignment = WD_ALIGN_PARAGRAPH.CENTER
    tight(par, after=10, lead=Pt(23))
    add_runs(par, h1, size=Pt(20))
    # Four authors go in a borderless 2x2 table, which is how the IEEE Word
    # template lays out more than two. Rendering by CSS class instead would
    # group all the names together and then all the affiliations.
    authors = tb.find_all("div", class_="author")
    cols = 2 if len(authors) > 1 else 1
    rows = (len(authors) + cols - 1) // cols
    at = doc.add_table(rows=rows, cols=cols)
    at.alignment = WD_TABLE_ALIGNMENT.CENTER
    at.autofit = False
    no_table_borders(at)
    set_cell_margins(at, left=0, right=0, top=0, bottom=0)
    cw = Inches(TEXT_WIDTH / cols)
    for col in at.columns:
        col.width = cw
    for idx, block_el in enumerate(authors):
        cell = at.rows[idx // cols].cells[idx % cols]
        cell.width = cw
        first = True
        for node in block_el.find_all("div"):
            klass = (node.get("class") or ["affil"])[0]
            par = cell.paragraphs[0] if first else cell.add_paragraph()
            first = False
            par.alignment = WD_ALIGN_PARAGRAPH.CENTER
            tight(par, after=0, lead=Pt(12))
            run = par.add_run(node.get_text().strip())
            run.font.name = MONO if klass == "mail" else SERIF
            run.font.size = {"name": Pt(11), "affil": Pt(9), "mail": Pt(8.5)}[klass]
            run.italic = (klass == "affil")
        if idx // cols < rows - 1:
            spacer = cell.add_paragraph()
            tight(spacer, after=0, lead=Pt(8))
    gap = doc.add_paragraph()
    tight(gap, after=0, lead=Pt(10))

    # --- body, two columns from here
    sec = doc.add_section(WD_SECTION.CONTINUOUS)
    set_columns(sec, 2)
    in_two = True
    table_i = 0

    body = soup.find("div", class_="body")
    for el in body.find_all(recursive=False):
        wide = "wide" in (el.get("class") or [])
        if wide and in_two:
            set_columns(doc.add_section(WD_SECTION.CONTINUOUS), 1)
            in_two = False
        elif not wide and not in_two:
            set_columns(doc.add_section(WD_SECTION.CONTINUOUS), 2)
            in_two = True

        if el.name == "p":
            emit_paragraph(doc, el)
        elif el.name == "h2":
            emit_heading(doc, el, 2)
        elif el.name == "h3":
            emit_heading(doc, el, 3)
        elif el.name == "ul":
            emit_list(doc, el)
        elif el.name == "figure":
            emit_figure(doc, el, TEXT_WIDTH - 0.16 if wide else COL_WIDTH)
        elif el.name == "div" and "tablewrap" in (el.get("class") or []):
            emit_table(doc, el, TABLE_WIDTHS[table_i])
            table_i += 1
        elif el.name == "ol" and "refs" in (el.get("class") or []):
            emit_references(doc, el)

    doc.save(OUT)
    print(f"wrote {OUT} ({OUT.stat().st_size / 1024:.0f} KB), {table_i} tables")


if __name__ == "__main__":
    main()
